"""Agrega el capítulo Módulo Inventario Tecnológico al Manual de Usuario existente."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

DOC_PATH = r"d:\rrhh\Manual_Usuario_Sistema_RRHH.docx"
COLOR_TEC  = RGBColor(0x4d, 0x7c, 0x8a)   # Azul petróleo Tecnología
COLOR_TEC_LIGHT = RGBColor(0xe8, 0xf2, 0xf5)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY  = RGBColor(0x4a, 0x60, 0x68)
COLOR_OK    = RGBColor(0x0b, 0x54, 0x47)
COLOR_OK_BG = RGBColor(0xd1, 0xfa, 0xe5)
COLOR_WARN_BG = RGBColor(0xfe, 0xf3, 0xc7)
COLOR_WARN  = RGBColor(0x92, 0x40, 0x0e)
COLOR_BORDER = RGBColor(0xcc, 0xdd, 0xe2)


def set_cell_bg(cell, hex_color: RGBColor):
    """Establece color de fondo a una celda."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), f"{hex_color[0]:02X}{hex_color[1]:02X}{hex_color[2]:02X}")
    tcPr.append(shd)


def set_cell_borders(cell, color="CCDDE2", size=4):
    """Agrega bordes a una celda."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), str(size))
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tcBorders.append(border)
    tcPr.append(tcBorders)


def add_heading(doc, text, level=1):
    """Agrega un título con el estilo del documento."""
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = COLOR_TEC
    return p


def add_paragraph(doc, text, bold=False, color=None, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = color
    return p


def add_colored_table_row(table, col1, col2, header=False):
    """Agrega una fila a una tabla de dos columnas."""
    row = table.add_row()
    row.cells[0].text = col1
    row.cells[1].text = col2
    if header:
        set_cell_bg(row.cells[0], COLOR_TEC)
        set_cell_bg(row.cells[1], COLOR_TEC)
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = COLOR_WHITE
                    run.font.bold = True
                    run.font.size = Pt(10)
    else:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
    return row


def add_callout(doc, title, body, bg_color, border_color):
    """Crea un bloque de nota/aviso con color de fondo."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_bg(cell, bg_color)
    set_cell_borders(cell, f"{border_color[0]:02X}{border_color[1]:02X}{border_color[2]:02X}", size=12)
    p = cell.paragraphs[0]
    run = p.add_run(f"{title}  ")
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(border_color[0], border_color[1], border_color[2])
    run2 = p.add_run(body)
    run2.font.size = Pt(10)
    doc.add_paragraph()


def add_step_list(doc, steps):
    """Agrega lista numerada de pasos."""
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.left_indent = Inches(0.25)
        if isinstance(step, tuple):
            run = p.add_run(step[0])
            run.font.bold = True
            run.font.size = Pt(10)
            run2 = p.add_run(step[1])
            run2.font.size = Pt(10)
        else:
            run = p.add_run(step)
            run.font.size = Pt(10)


def add_bullet_list(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25)
        if isinstance(item, tuple):
            run = p.add_run(item[0])
            run.font.bold = True
            run.font.size = Pt(10)
            run2 = p.add_run(item[1])
            run2.font.size = Pt(10)
        else:
            run = p.add_run(item)
            run.font.size = Pt(10)


# ─────────────────────────────────────────────
doc = Document(DOC_PATH)

# Salto de página antes del nuevo capítulo
doc.add_page_break()

# ══════════════════════════════════════════════
# CAPÍTULO: MÓDULO INVENTARIO TECNOLÓGICO
# ══════════════════════════════════════════════
add_heading(doc, "MÓDULO INVENTARIO TECNOLÓGICO", level=1)

p = doc.add_paragraph(
    "Este módulo permite a la Dirección de Tecnologías llevar el control completo del inventario "
    "de equipos tecnológicos institucionales: estado físico, custodia (quién tiene cada equipo), "
    "mantenimiento preventivo anual y trazabilidad de piezas y repuestos cambiados."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

# ── ACCESO Y ROLES ──
add_heading(doc, "Acceso y roles", level=2)

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(4.0)

# Encabezado
add_colored_table_row(table, "ROL", "PERMISOS", header=True)

r1 = add_colored_table_row(table, "TECNOLOGÍA", "")
r1.cells[1].text = (
    "• Acceso completo a todos los equipos\n"
    "• Gestionar custodia (asignar / devolver)\n"
    "• Registrar y aprobar mantenimientos\n"
    "• Administrar piezas y repuestos\n"
    "• Importar inventario desde CSV\n"
    "• Configurar tipos de equipo y checklist"
)
for p in r1.cells[1].paragraphs:
    for run in p.runs:
        run.font.size = Pt(10)

r2 = add_colored_table_row(table, "ADMINISTRADOR", "Mismos permisos que el rol Tecnología. Acceso desde el launcher institucional.")
for p in r2.cells[1].paragraphs:
    for run in p.runs:
        run.font.size = Pt(10)

doc.add_paragraph()
add_callout(doc,
    "ℹ️  Importante: ",
    "Para asignar el rol Tecnología a un empleado, ir a Administración → Roles.",
    COLOR_TEC_LIGHT, COLOR_TEC
)

# ── MENÚ DE NAVEGACIÓN ──
add_heading(doc, "Menú de navegación", level=2)

table2 = doc.add_table(rows=1, cols=2)
table2.style = 'Table Grid'
table2.alignment = WD_TABLE_ALIGNMENT.LEFT
table2.columns[0].width = Inches(2.8)
table2.columns[1].width = Inches(3.7)
add_colored_table_row(table2, "Opción del menú", "Descripción", header=True)

menu_items = [
    ("Inventario de Equipos",             "Ver, registrar, asignar y gestionar todos los equipos tecnológicos"),
    ("Mantenimiento",                      "Registrar mantenimientos preventivos anuales por equipo"),
    ("Piezas y Repuestos",                 "Catálogo de piezas entregadas por Bienes e historial de instalación"),
    ("Tipos de Equipo",                    "Configurar las categorías de equipos (Laptop, Impresora, etc.)"),
    ("Actividades de Mantenimiento",       "Configurar el checklist del formulario de mantenimiento"),
]
for name, desc in menu_items:
    add_colored_table_row(table2, name, desc)

doc.add_paragraph()

# ══════════════════════════════════════════════
# GESTIÓN DE EQUIPOS
# ══════════════════════════════════════════════
add_heading(doc, "Gestión de equipos", level=2)

add_heading(doc, "Panel principal — Inventario", level=3)
p = doc.add_paragraph(
    "Al ingresar a Inventario de Equipos se muestran seis tarjetas de resumen en la parte superior. "
    "Hacer clic en cualquier tarjeta filtra automáticamente la tabla:"
)
p.runs[0].font.size = Pt(11)

table3 = doc.add_table(rows=1, cols=2)
table3.style = 'Table Grid'
table3.columns[0].width = Inches(2.0)
table3.columns[1].width = Inches(4.5)
add_colored_table_row(table3, "Tarjeta", "Descripción", header=True)
estados = [
    ("DISPONIBLE",        "Equipos sin asignar, listos para custodia"),
    ("ASIGNADO",          "Equipos que tienen un custodio activo"),
    ("DAÑADO",            "Equipos devueltos por daño; pendientes de revisión o baja"),
    ("DE BAJA",           "Equipos dados de baja; no se reasignan"),
    ("Vida útil vencida", "Equipos cuya vida útil (en años) ya expiró — requieren atención"),
    ("Total",             "Conteo total del inventario"),
]
for e, d in estados:
    add_colored_table_row(table3, e, d)
doc.add_paragraph()

add_heading(doc, "Registrar un nuevo equipo", level=3)
add_step_list(doc, [
    ("Clic en Nuevo Equipo ", "(botón esquina superior derecha)."),
    ("Completar los campos: ", "Código del bien, Tipo de equipo, Marca, Modelo, Serie, Descripción, Condición (Bueno / Regular / Malo), Fecha de ingreso y Vida útil (años)."),
    ("Clic en Guardar. ", "El equipo queda en estado DISPONIBLE."),
])
doc.add_paragraph()

add_heading(doc, "Importar inventario desde CSV", level=3)
p = doc.add_paragraph("Para cargar múltiples equipos a la vez, usar el botón Importar CSV. Columnas requeridas:")
p.runs[0].font.size = Pt(11)

cols_p = doc.add_paragraph()
cols_p.paragraph_format.left_indent = Inches(0.3)
run = cols_p.add_run(
    "codigo_bien  |  tipo_equipo  |  marca  |  modelo  |  descripcion  |  serie  |  estado  |  fecha_ingreso  |  vida_util_anios"
)
run.font.name = 'Courier New'
run.font.size = Pt(9)

add_callout(doc,
    "📋  Nota: ",
    "El sistema acepta archivos separados por coma (,) o punto y coma (;). "
    "Las fechas pueden estar en formato DD/MM/AAAA o AAAA-MM-DD. "
    "Si hay un error en cualquier fila, no se importa nada y se muestran todos los errores juntos.",
    COLOR_TEC_LIGHT, COLOR_TEC
)

# ══════════════════════════════════════════════
# CUSTODIA
# ══════════════════════════════════════════════
add_heading(doc, "Custodia de equipos", level=2)

add_heading(doc, "Asignar un equipo", level=3)
p = doc.add_paragraph("Solo se puede asignar un equipo con estado DISPONIBLE.")
p.runs[0].font.size = Pt(11)
add_step_list(doc, [
    ("En la fila del equipo, ", "clic en el botón Asignar."),
    ("En el modal, ", "buscar al empleado escribiendo su nombre o cédula."),
    ("Seleccionar el empleado ", "y registrar la fecha de asignación."),
    ("Clic en Guardar asignación. ", "El equipo pasa a estado ASIGNADO y muestra el nombre del custodio."),
])
doc.add_paragraph()

add_heading(doc, "Devolver un equipo", level=3)
p = doc.add_paragraph("Solo se puede devolver un equipo con estado ASIGNADO.")
p.runs[0].font.size = Pt(11)
add_step_list(doc, [
    ("Clic en Devolver ", "en la fila del equipo."),
    ("Registrar la fecha de devolución ", "(puede ser retroactiva)."),
    ("Seleccionar el motivo: ", "Reasignación / Salida del empleado / Daño / Otro."),
    ("Agregar observación ", "si es necesario y clic en Guardar."),
])
doc.add_paragraph()

add_callout(doc,
    "⚠️  Motivo Daño: ",
    "Cuando se selecciona motivo Daño, el equipo cambia automáticamente al estado DAÑADO "
    "en lugar de volver a Disponible.",
    COLOR_WARN_BG, COLOR_WARN
)

add_heading(doc, "Historial de custodia", level=3)
p = doc.add_paragraph(
    "El botón Historial (columna de acciones) abre una línea de tiempo completa de todas las "
    "asignaciones del equipo: custodio, fechas de asignación y devolución, y motivo de cada devolución. "
    "La asignación activa se muestra con un punto verde; las históricas con punto gris."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

# ══════════════════════════════════════════════
# MANTENIMIENTO
# ══════════════════════════════════════════════
add_heading(doc, "Mantenimiento", level=2)

p = doc.add_paragraph(
    "El mantenimiento preventivo se realiza una vez al año por equipo. "
    "La vista tiene dos pestañas: Pendientes {año} (equipos sin mantenimiento en el año actual) "
    "y Realizados {año} (equipos con mantenimiento ya registrado). "
    "El gráfico de dona muestra el porcentaje de avance anual."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

add_heading(doc, "Mantenimiento interno (individual)", level=3)
p = doc.add_paragraph("Lo ejecuta un técnico de la Dirección de Tecnología directamente en el equipo.")
p.runs[0].font.size = Pt(11)
add_step_list(doc, [
    ("En la pestaña Pendientes, ", "ubicar el equipo y clic en Registrar mantenimiento."),
    ("Ingresar la fecha, ", "hora de inicio y hora de fin del trabajo."),
    ("Completar el checklist ", "marcando Sí / No para cada actividad (10 ítems). El técnico y el custodio se capturan automáticamente."),
    ("Si se instalaron piezas, ", "agregarlas en la sección Piezas Cambiadas."),
    ("Clic en Guardar. ", "Se genera el acta de mantenimiento en PDF con checklist y firmas."),
    ("Imprimir el acta, obtener firmas físicas ", "y subir el PDF firmado con el botón Subir Firmado."),
])
doc.add_paragraph()

add_heading(doc, "Checklist de mantenimiento (10 ítems estándar)", level=3)
checklist_items = [
    "Ingreso al equipo",
    "Limpieza interna del equipo",
    "Limpieza externa del equipo",
    "Borrado de archivos temporales",
    "Ingreso al equipo por la IP",
    "Actualización del antivirus",
    "Formateo del equipo",
    "Respaldo carpeta Escritorio",
    "Respaldo carpeta Mis documentos",
    "Respaldo correo electrónico institucional (PST)",
]
add_bullet_list(doc, checklist_items)
doc.add_paragraph()

add_heading(doc, "Mantenimiento externo (por proveedor, en lote)", level=3)
p = doc.add_paragraph(
    "Para equipos que atiende un proveedor externo bajo una orden de compra "
    "(por ejemplo infraestructura de Data Center o equipos de videovigilancia). "
    "El proveedor atiende todos los equipos de una categoría en una sola intervención."
)
p.runs[0].font.size = Pt(11)
add_step_list(doc, [
    ("En la pestaña Pendientes, ", "clic en Registrar mantenimiento externo."),
    ("Seleccionar: ", "Categoría de equipo, Fecha, Proveedor, Proceso de contratación y N.° de orden de compra."),
    ("Clic en Guardar. ", "El sistema registra el mantenimiento para todos los equipos de esa categoría que no lo tenían en el año y genera el acta grupal en PDF."),
])
doc.add_paragraph()

add_callout(doc,
    "ℹ️  En la pestaña Realizados: ",
    "Los mantenimientos externos se agrupan en una sola tarjeta con etiqueta EXTERNO "
    "mostrando el proveedor, proceso y cantidad de equipos. "
    "Los internos se muestran uno por uno con etiqueta INTERNO.",
    COLOR_TEC_LIGHT, COLOR_TEC
)

# ══════════════════════════════════════════════
# PIEZAS Y REPUESTOS
# ══════════════════════════════════════════════
add_heading(doc, "Piezas y repuestos", level=2)

p = doc.add_paragraph(
    "Registra las piezas (discos, memorias RAM, fuentes de poder, etc.) que la Unidad de Bienes "
    "entrega a Tecnología para instalación en equipos. Tecnología no tiene bodega propia — "
    "las piezas llegan ya codificadas desde Bienes."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

add_heading(doc, "Ciclo de vida de una pieza", level=3)
p = doc.add_paragraph(
    "DISPONIBLE (entregada por Bienes)  →  INSTALADA (en un equipo)  →  "
    "DISPONIBLE (retirada, reutilizable)  o  DE BAJA (retirada por daño)"
)
p.runs[0].font.size = Pt(10)
p.runs[0].font.color.rgb = COLOR_TEC
doc.add_paragraph()

add_heading(doc, "Registrar una nueva pieza", level=3)
add_step_list(doc, [
    ("En Piezas y Repuestos, clic en Nueva Pieza."),
    ("Completar: ", "Descripción (obligatoria), Código y Serie (opcionales) y Fecha de entrega de Bienes (obligatoria — fecha en que Bienes entregó la pieza a Tecnología)."),
    ("Clic en Guardar. ", "La pieza queda en estado DISPONIBLE."),
])
doc.add_paragraph()

add_heading(doc, "Instalar y retirar piezas", level=3)
p = doc.add_paragraph("Hay dos formas de instalar una pieza:")
p.runs[0].font.size = Pt(11)
add_bullet_list(doc, [
    ("Desde Piezas y Repuestos: ", "buscar la pieza → botón Instalar → buscar el equipo → Guardar."),
    ("Desde el formulario de Mantenimiento: ", 'sección "Piezas Cambiadas", buscar piezas disponibles o crear una nueva al vuelo. Al guardar el mantenimiento, las piezas se instalan automáticamente.'),
])
doc.add_paragraph()

p = doc.add_paragraph(
    "Para retirar una pieza instalada, usar el botón Retirar en la fila de la pieza. "
    "La pieza vuelve a estado DISPONIBLE o se puede marcar como DE BAJA si está dañada."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

add_heading(doc, "Historial de una pieza", level=3)
p = doc.add_paragraph(
    "El botón Historial muestra la línea de tiempo de todos los equipos en los que ha estado "
    "instalada la pieza, con fechas y el mantenimiento vinculado (si aplica)."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

# ══════════════════════════════════════════════
# CONFIGURACIÓN
# ══════════════════════════════════════════════
add_heading(doc, "Configuración", level=2)

add_heading(doc, "Tipos de equipo", level=3)
p = doc.add_paragraph(
    "Desde Tipos de Equipo se puede agregar, editar o desactivar las categorías del inventario. "
    "Ejemplos predeterminados: Computador de Escritorio, Laptop, Impresora, Monitor, Escáner, Proyector."
)
p.runs[0].font.size = Pt(11)
doc.add_paragraph()

add_heading(doc, "Actividades de mantenimiento", level=3)
p = doc.add_paragraph(
    "Desde Actividades de Mantenimiento se puede gestionar el checklist que aparece en cada "
    "mantenimiento interno: cambiar el nombre de un ítem, agregar nuevos o desactivar los que ya "
    "no aplican. El orden se ajusta con el campo Orden."
)
p.runs[0].font.size = Pt(11)

# ─── GUARDAR ───
doc.save(DOC_PATH)
print(f"✅ Capítulo 'Módulo Inventario Tecnológico' agregado exitosamente a:\n   {DOC_PATH}")
